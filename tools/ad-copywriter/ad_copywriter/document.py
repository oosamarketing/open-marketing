"""The copy document: one brand + one platform + themed sets of options.

This is the interchange format between the agent, the scripts, the exports,
the previews and any app that embeds this library. See ``data/copy_document.schema.json``.
"""

from __future__ import annotations

import datetime as _dt
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .spec import PlatformSpec, get_platform, platform_keys

SCHEMA_VERSION = 1


def _norm_media_paths(raw: Any) -> list[dict[str, str]]:
    """Accept ["a.png", ...] or [{"square": "a.png", "wide": "b.png"}, ...]; strings become {"square": s}."""
    out: list[dict[str, str]] = []
    for item in raw or []:
        if isinstance(item, str) and item:
            out.append({"square": item})
        elif isinstance(item, dict):
            m = {k: str(v) for k, v in item.items() if k in ("square", "wide") and v}
            if m:
                out.append(m)
    return out


@dataclass
class CopySet:
    theme: str
    fields: dict[str, list[str]] = field(default_factory=dict)
    landing_page_url: str = ""
    display_url: str = ""
    media_folder_link: str = ""
    media_path: str = ""
    media_paths: list[dict[str, str]] = field(default_factory=list)  # [{"square": path, "wide": path}, ...] per variation
    cta: str = "Shop now"
    audiences: dict[str, str] = field(default_factory=dict)
    notes: str = ""

    def options(self, key: str) -> list[str]:
        return [o for o in self.fields.get(key, []) if o is not None]

    def to_dict(self) -> dict[str, Any]:
        return {
            "theme": self.theme,
            "landing_page_url": self.landing_page_url or None,
            "display_url": self.display_url or None,
            "media_folder_link": self.media_folder_link or None,
            "media_path": self.media_path or None,
            "media_paths": [dict(m) for m in self.media_paths] or None,
            "cta": self.cta or None,
            "audiences": dict(self.audiences),
            "fields": {k: list(v) for k, v in self.fields.items()},
            "notes": self.notes or None,
        }

    @classmethod
    def from_dict(cls, d: dict[str, Any]) -> "CopySet":
        return cls(
            theme=str(d.get("theme") or "").strip(),
            fields={k: [str(x) for x in (v or [])] for k, v in (d.get("fields") or {}).items()},
            landing_page_url=d.get("landing_page_url") or "",
            display_url=d.get("display_url") or "",
            media_folder_link=d.get("media_folder_link") or "",
            media_path=d.get("media_path") or "",
            media_paths=_norm_media_paths(d.get("media_paths")),
            cta=d.get("cta") or "Shop now",
            audiences={k: (v or "") for k, v in (d.get("audiences") or {}).items()},
            notes=d.get("notes") or "",
        )


@dataclass
class CopyDocument:
    platform: str
    brand: dict[str, Any]
    sets: list[CopySet] = field(default_factory=list)
    validation: dict[str, Any] | None = None
    meta: dict[str, Any] = field(default_factory=dict)
    schema_version: int = SCHEMA_VERSION

    # -- convenience -------------------------------------------------------
    @property
    def spec(self) -> PlatformSpec:
        return get_platform(self.platform)

    @property
    def brand_name(self) -> str:
        return str(self.brand.get("name") or "Brand")

    @property
    def themes(self) -> list[str]:
        return [s.theme for s in self.sets]

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "platform": self.platform,
            "brand": dict(self.brand),
            "sets": [s.to_dict() for s in self.sets],
            "validation": self.validation,
            "meta": dict(self.meta),
        }

    @classmethod
    def from_dict(cls, d: dict[str, Any]) -> "CopyDocument":
        if not isinstance(d, dict):
            raise ValueError("copy document must be a JSON object")
        platform = str(d.get("platform") or "").strip()
        if platform not in platform_keys():
            raise ValueError(f"unknown platform {platform!r}; expected one of {platform_keys()}")
        brand = d.get("brand") or {}
        if isinstance(brand, str):
            brand = {"name": brand}
        if not brand.get("name"):
            raise ValueError("brand.name is required")
        sets_raw = d.get("sets") or []
        if not isinstance(sets_raw, list) or not sets_raw:
            raise ValueError("sets must be a non-empty list")
        return cls(
            platform=platform,
            brand=dict(brand),
            sets=[CopySet.from_dict(s) for s in sets_raw],
            validation=d.get("validation"),
            meta=dict(d.get("meta") or {}),
            schema_version=int(d.get("schema_version") or SCHEMA_VERSION),
        )


# -- IO ----------------------------------------------------------------------

def load_document(src: str | Path | dict[str, Any]) -> CopyDocument:
    """Load from a path, a JSON string, or an already-parsed dict."""
    if isinstance(src, dict):
        return CopyDocument.from_dict(src)
    if isinstance(src, (str, Path)) and not str(src).lstrip().startswith("{"):
        with open(src, encoding="utf-8") as fh:
            return CopyDocument.from_dict(json.load(fh))
    return CopyDocument.from_dict(json.loads(str(src)))


def dump_document(doc: CopyDocument, path: str | Path | None = None) -> str:
    text = json.dumps(doc.to_dict(), indent=2, ensure_ascii=False) + "\n"
    if path is not None:
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(text, encoding="utf-8")
    return text


def new_document(
    platform: str,
    brand_name: str,
    themes: list[str] | None = None,
    *,
    website: str = "",
    identity: str = "",
    logo_path: str | None = None,
    generated_by: str = "agent",
    **brand_extra: Any,
) -> CopyDocument:
    """Scaffold an empty document with the right field keys for the platform."""
    spec = get_platform(platform)
    brand = {
        "name": brand_name,
        "website": website or None,
        "page_name": brand_extra.pop("page_name", None) or brand_name,
        "logo_path": logo_path,
        "primary_color": brand_extra.pop("primary_color", None),
        "secondary_color": brand_extra.pop("secondary_color", None),
        "font_family": brand_extra.pop("font_family", None),
        "identity": identity or None,
    }
    sets = [
        CopySet(theme=t, fields={f.key: [] for f in spec.fields if not f.optional})
        for t in (["Untitled theme"] if themes is None else themes)
    ]
    return CopyDocument(
        platform=platform,
        brand=brand,
        sets=sets,
        meta={
            "generated_by": generated_by,
            "model": None,
            "created_at": _dt.datetime.now(_dt.timezone.utc).isoformat(timespec="seconds"),
        },
    )


# -- xlsx importers ------------------------------------------------------------

def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def from_legacy_xlsx(path: str | Path, platform: str, brand_name: str, **brand_extra: Any) -> CopyDocument:
    """Import the legacy layout: one sheet per theme, header row of field labels
    (``Primary Text | Headline | Description``), one option per row.
    """
    from openpyxl import load_workbook

    spec = get_platform(platform)
    wb = load_workbook(path, read_only=True, data_only=True)
    doc = new_document(platform, brand_name, themes=[], **brand_extra)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [_cell_text(h) for h in rows[0]]
        columns: dict[int, str] = {}
        for idx, h in enumerate(header):
            f = spec.legacy_field_for_column(h) if h else None
            if f:
                columns[idx] = f.key
        if not columns:
            continue
        fields: dict[str, list[str]] = {k: [] for k in columns.values()}
        for row in rows[1:]:
            for idx, key in columns.items():
                if idx < len(row):
                    text = _cell_text(row[idx])
                    if text:
                        fields[key].append(text)
        doc.sets.append(CopySet(theme=ws.title, fields=fields))
    if not doc.sets:
        raise ValueError(f"no sheets with recognisable columns ({spec.legacy_columns}) in {path}")
    return doc


def from_assets_xlsx(path: str | Path, platform: str, brand_name: str, **brand_extra: Any) -> CopyDocument:
    """Import the oosamarketing.com tool layout: ``Asset Type | Content | Character Count``
    with rows like ``Headline 3``. Each sheet becomes a set.
    """
    from openpyxl import load_workbook

    spec = get_platform(platform)
    labels = {f.label.lower(): f.key for f in spec.fields}
    row_keys = {r.label.lower(): r.key for r in spec.rows}
    opt_keys = {r.label.lower(): r.key for r in spec.optional_rows}
    wb = load_workbook(path, read_only=True, data_only=True)
    doc = new_document(platform, brand_name, themes=[], **brand_extra)
    for ws in wb.worksheets:
        cs = CopySet(theme=ws.title, fields={})
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not _cell_text(row[0]):
                continue
            asset = _cell_text(row[0])
            content = _cell_text(row[1]) if len(row) > 1 else ""
            base = asset.rsplit(" ", 1)[0].lower() if asset[-1:].isdigit() else asset.lower()
            if base in labels:
                if content:
                    cs.fields.setdefault(labels[base], []).append(content)
            elif asset.lower() in row_keys:
                setattr(cs, row_keys[asset.lower()], content)
            elif asset.lower() in opt_keys:
                if content:
                    cs.audiences[opt_keys[asset.lower()]] = content
        if cs.fields:
            doc.sets.append(cs)
    if not doc.sets:
        raise ValueError(f"no recognisable asset rows in {path}")
    return doc
