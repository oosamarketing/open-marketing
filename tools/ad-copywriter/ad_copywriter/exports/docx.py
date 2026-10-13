"""docx writer (python-docx) — a copy deck document mirroring the PDF deck layout."""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

from ..document import CopyDocument
from ..textutil import count_chars, split_blocks

FONT = "Montserrat"
GOLD = "F0E2B6"
MUTED = "6B6B6B"


def _set_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    from docx.shared import Pt, RGBColor

    run.font.name = FONT
    # East-Asian fallback so Word honours the font name everywhere
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    from docx.oxml.ns import qn

    rfonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _rule(doc) -> None:
    """A gold horizontal rule, done as a bottom border on an empty paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GOLD)
    pbdr.append(bottom)
    pPr.append(pbdr)


def write_docx(doc: CopyDocument, path: str | Path, *, logo: str | Path | None = None,
               title: str | None = None, subtitle: str | None = None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    spec = doc.spec
    d = Document()
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            d.styles[style_name].font.name = FONT
        except KeyError:
            pass

    logo_path = Path(logo) if logo else None
    if logo_path is None and doc.brand.get("logo_path"):
        candidate = Path(str(doc.brand["logo_path"]))
        if candidate.exists():
            logo_path = candidate
    if logo_path and logo_path.exists():
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo_path), width=Inches(1.2))

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(title or spec.deck_title), size=24, bold=True)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = subtitle or f"{doc.brand_name} · {spec.deck_subtitle} · {_dt.date.today():%B %d, %Y}"
    _set_font(p.add_run(sub), size=11, color=MUTED)
    _rule(d)

    for si, cs in enumerate(doc.sets):
        if si:
            d.add_page_break()
        h = d.add_heading(cs.theme, level=1)
        for r in h.runs:
            _set_font(r)
        if cs.landing_page_url or cs.display_url:
            p = d.add_paragraph()
            _set_font(p.add_run(" · ".join(x for x in (cs.display_url, cs.landing_page_url) if x)),
                      size=9, color=MUTED)
        for f in spec.fields:
            options = cs.options(f.key)
            if not options:
                continue
            h = d.add_heading(f.plural, level=2)
            for r in h.runs:
                _set_font(r)
            for i, text in enumerate(options):
                if f.multiline:
                    blocks = split_blocks(text)
                    p = d.add_paragraph(style="List Number")
                    for bi, block in enumerate(blocks):
                        if bi:
                            p.add_run("\n...\n")
                        _set_font(p.add_run(block), size=10.5)
                else:
                    p = d.add_paragraph(style="List Number")
                    _set_font(p.add_run(text), size=10.5)
                    _set_font(p.add_run(f"   ({count_chars(text)}/{f.max_chars})"), size=8, color=MUTED)
                p.paragraph_format.space_after = Pt(6)
        _rule(d)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path
