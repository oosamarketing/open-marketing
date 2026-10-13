"""xlsx writers (openpyxl)."""

from __future__ import annotations

from pathlib import Path

from ..document import CopyDocument, CopySet
from ..spec import PlatformSpec
from ..textutil import count_chars, safe_sheet_name

HEADER_FILL = "F0E2B6"  # OOSA gold
HEADER_FONT = "1B1B1B"


def _style_header(ws, ncols: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color=HEADER_FONT)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _wrap_all(ws) -> None:
    from openpyxl.styles import Alignment

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def assets_rows(cs: CopySet, spec: PlatformSpec) -> list[tuple[str, str, str]]:
    """The fixed row list of the site tool: numbered asset rows, URL rows, optional audience rows."""
    rows: list[tuple[str, str, str]] = []
    for f in spec.fields:
        options = cs.options(f.key)
        if f.optional and not options:
            continue
        for i in range(f.count):
            text = options[i] if i < len(options) else ""
            count = f"{count_chars(text)} / {f.max_chars}" if text else ""
            rows.append((f.row_label(i), text, count))
    for r in spec.rows:
        value = getattr(cs, r.key, "") or ""
        rows.append((r.label, value, "—" if value else ""))
    for r in spec.optional_rows:
        value = cs.audiences.get(r.key, "") or ""
        if value:
            rows.append((r.label, value, "—"))
    return rows


def write_assets_xlsx(doc: CopyDocument, path: str | Path) -> Path:
    from openpyxl import Workbook

    spec = doc.spec
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for cs in doc.sets:
        title = spec.sheet_name if len(doc.sets) == 1 else f"{cs.theme[:20]} Assets"
        ws = wb.create_sheet(safe_sheet_name(title, used))
        ws.append(["Asset Type", "Content", "Character Count"])
        for row in assets_rows(cs, spec):
            ws.append(list(row))
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 90
        ws.column_dimensions["C"].width = 16
        _style_header(ws, 3)
        _wrap_all(ws)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_legacy_xlsx(doc: CopyDocument, path: str | Path) -> Path:
    """One sheet per theme; columns are the platform's legacy field labels; options down the rows."""
    from openpyxl import Workbook

    spec = doc.spec
    fields = [f for f in spec.fields if f.legacy_column]
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for cs in doc.sets:
        ws = wb.create_sheet(safe_sheet_name(cs.theme, used))
        ws.append([f.legacy_column for f in fields])
        columns = [cs.options(f.key) for f in fields]
        depth = max((len(c) for c in columns), default=0)
        for i in range(depth):
            ws.append([c[i] if i < len(c) else None for c in columns])
        for idx, f in enumerate(fields):
            ws.column_dimensions[chr(ord("A") + idx)].width = 70 if f.multiline else 40
        _style_header(ws, len(fields))
        _wrap_all(ws)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
